import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.app import app
from src.main import main
from src.services.pipeline import PipelineRun
from src.utils.date_utils import build_report_period

client = TestClient(app)


def wait_for(run_id, timeout=30):
    deadline = time.time() + timeout
    while time.time() < deadline:
        state = client.get(f"/api/runs/{run_id}").json()
        if state["status"] in ("succeeded", "failed"):
            return state
        time.sleep(0.1)
    raise AssertionError("run did not finish")


def test_pipeline_produces_all_outputs(tmp_path):
    period = build_report_period("date", "22-09-2026")
    state = PipelineRun(period, data_source="demo", runs_dir=tmp_path).execute()
    assert state["status"] == "succeeded"
    assert [s["status"] for s in state["stages"]] == ["done", "done", "done", "done", "skipped", "done", "done"]
    run_dir = tmp_path / state["id"]
    workbook = load_workbook(run_dir / state["artifacts"]["excel"])
    assert workbook.sheetnames == ["Summary", "Breach periods", "Hourly peaks", "Graphs", "Accuracy checks"]
    assert (run_dir / "images" / "Hyderabad.png").exists()
    assert state["result"]["counts"]["sites"] == 7


def test_api_run_lifecycle():
    created = client.post("/api/runs", json={"mode": "date", "date": "2026-09-21"})
    assert created.status_code == 202
    state = wait_for(created.json()["id"])
    assert state["status"] == "succeeded"
    assert state["period"]["label"] == "21-09-2026"
    assert client.get(f"/api/runs/{state['id']}/files/excel").status_code == 200
    assert "Network utilization update" in client.get(f"/api/runs/{state['id']}/files/message").text
    assert any(r["id"] == state["id"] for r in client.get("/api/runs").json())
    assert client.get("/api/runs/latest").json()["id"] == state["id"]


def test_api_rejects_bad_requests():
    assert client.post("/api/runs", json={"mode": "date", "date": "2099-01-01"}).status_code == 422
    assert client.post("/api/runs", json={"mode": "sometime"}).status_code == 422
    assert client.post("/api/runs", json={"mode": "yesterday", "source": "snmp"}).status_code == 422
    assert client.get("/api/runs/does-not-exist").status_code == 404
    assert client.get("/api/runs/..%2F..%2Fetc/files/excel").status_code == 404


def test_dashboard_is_served():
    page = client.get("/")
    assert page.status_code == 200 and "Network utilization" in page.text


def test_cli_non_interactive():
    assert main(["--mode", "date", "--date", "20-09-2026"]) == 0
    assert main(["--mode", "date", "--date", "not-a-date"]) == 2
