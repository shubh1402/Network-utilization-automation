"""Excel workbook the ops team receives: summary, episodes, hourly heatmap, graphs, checks."""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.config import THRESHOLDS
from src.messages.formatting import period_title
from src.services.utilization_service import build_headline

HEADER_FILL = PatternFill("solid", fgColor="1D2B36")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILLS = {
    "critical": PatternFill("solid", fgColor="F8D0D1"),
    "high": PatternFill("solid", fgColor="FBE0CC"),
    "elevated": PatternFill("solid", fgColor="FBF0C6"),
    "normal": PatternFill("solid", fgColor="DDEFE9"),
    "no-data": PatternFill("solid", fgColor="E6E6E6"),
}
THIN = Side(style="thin", color="D5DADF")
BORDER = Border(bottom=THIN)


def _header(ws, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _widths(ws, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _ts(iso_value: str | None):
    return datetime.fromisoformat(iso_value) if iso_value else None


def _summary_sheet(wb: Workbook, label: str, site_data: dict, meta: dict) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Network utilization report: {period_title(label)}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = build_headline(site_data)
    ws["A2"].font = Font(italic=True)
    ws["A3"] = f"Generated {meta.get('generated_at', '')} | Data source: {meta.get('data_source', '')}"
    ws["A3"].font = Font(color="6B7B88", size=9)

    headers = ["Site", "Link", "Status", "Peak %", "Peak time", "Average %", "P95 %"] + [
        f"Minutes >{t}%" for t in THRESHOLDS
    ] + ["Samples"]
    _header(ws, 5, headers)
    row = 6
    for site, links in site_data.items():
        for link_type, link in links.items():
            values = [
                site, link_type, link["status"], link["peak"], _ts(link["peak_time"]),
                link["average"], link["p95"],
                *[link[f"above_{t}_events"] for t in THRESHOLDS], link["total_base_events"],
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = BORDER
                if col == 3:
                    cell.fill = STATUS_FILLS.get(link["status"], STATUS_FILLS["no-data"])
                if col == 5 and value is not None:
                    cell.number_format = "dd-mm-yyyy hh:mm"
                if col in (4, 6, 7):
                    cell.number_format = "0.0"
            row += 1
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(headers))}{row - 1}"
    _widths(ws, [14, 11, 11, 9, 17, 11, 9, 12, 12, 12, 10])


def _episodes_sheet(wb: Workbook, site_data: dict) -> None:
    ws = wb.create_sheet("Breach periods")
    headers = ["Site", "Link", "Start", "End", "Duration (min)", "Minutes >70%", "Peak %", "Peak time", "Highest level"]
    _header(ws, 1, headers)
    row = 2
    for site, links in site_data.items():
        for link_type, link in links.items():
            for ep in link.get("episodes", []):
                values = [site, link_type, _ts(ep["start"]), _ts(ep["end"]), ep["duration"], ep["minutes"],
                          ep["peak"], _ts(ep["peak_time"]), f">{ep['level']}%"]
                for col, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.border = BORDER
                    if isinstance(value, datetime):
                        cell.number_format = "dd-mm-yyyy hh:mm"
                level_fill = {90: "critical", 80: "high", 70: "elevated"}[ep["level"]]
                ws.cell(row=row, column=9).fill = STATUS_FILLS[level_fill]
                row += 1
    if row == 2:
        ws.cell(row=2, column=1, value="No link went above 70% in this period.")
    ws.freeze_panes = "A2"
    _widths(ws, [14, 11, 17, 17, 14, 13, 9, 17, 13])


def _hourly_sheet(wb: Workbook, site_data: dict, series: dict, start: datetime, step: int) -> None:
    """Hour x link grid of peak utilization, coloured like a heatmap."""
    ws = wb.create_sheet("Hourly peaks")
    columns = [(site, link) for site, links in site_data.items() for link in links]
    _header(ws, 1, ["Hour"] + [f"{site}\n{link}" for site, link in columns])
    per_hour = 60 // step
    n_hours = max((len(series.get(s, {}).get(lk, [])) + per_hour - 1) // per_hour for s, lk in columns) if columns else 0

    hourly: dict[int, dict[int, float]] = defaultdict(dict)
    for col, (site, link) in enumerate(columns):
        values = series.get(site, {}).get(link, [])
        for hour in range(n_hours):
            chunk = [v for v in values[hour * per_hour:(hour + 1) * per_hour] if v is not None]
            if chunk:
                hourly[hour][col] = max(chunk)


    for hour in range(n_hours):
        cell = ws.cell(row=hour + 2, column=1, value=start + timedelta(hours=hour))
        cell.number_format = "dd-mm hh:mm"
        for col in range(len(columns)):
            value = hourly[hour].get(col)
            c = ws.cell(row=hour + 2, column=col + 2, value=value)
            c.number_format = "0"
            c.alignment = Alignment(horizontal="center")

    if n_hours and columns:
        rng = f"B2:{get_column_letter(len(columns) + 1)}{n_hours + 1}"
        ws.conditional_formatting.add(
            rng,
            ColorScaleRule(start_type="num", start_value=0, start_color="E8F3F1",
                           mid_type="num", mid_value=70, mid_color="F6E27A",
                           end_type="num", end_value=100, end_color="D13438"),
        )
    ws.freeze_panes = "B2"
    _widths(ws, [13] + [11] * len(columns))


def _graphs_sheet(wb: Workbook, screenshots_dir: Path | None) -> None:
    ws = wb.create_sheet("Graphs")
    images = sorted(Path(screenshots_dir).glob("*.png")) if screenshots_dir and Path(screenshots_dir).exists() else []
    if not images:
        ws["A1"] = "No graphs were captured for this run."
        return
    row = 1
    for image_path in images:
        ws.cell(row=row, column=1, value=image_path.stem.replace("_", " ")).font = Font(bold=True)
        image = XLImage(str(image_path))
        scale = 820 / image.width
        image.width, image.height = int(image.width * scale), int(image.height * scale)
        ws.add_image(image, f"A{row + 1}")
        row += int(image.height / 20) + 3


def _checks_sheet(wb: Workbook, checks: list[dict]) -> None:
    ws = wb.create_sheet("Accuracy checks")
    _header(ws, 1, ["Check", "Result", "Detail"])
    for row, check in enumerate(checks, start=2):
        result = "Pass" if check["passed"] else ("Warning" if check["severity"] == "warning" else "Fail")
        ws.cell(row=row, column=1, value=check["name"])
        cell = ws.cell(row=row, column=2, value=result)
        cell.fill = STATUS_FILLS["normal" if check["passed"] else ("elevated" if result == "Warning" else "critical")]
        ws.cell(row=row, column=3, value=check["detail"]).alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, [28, 10, 90])


def save_excel_report(
    output_dir: Path,
    report_label: str,
    site_data: dict,
    screenshots_dir: Path | None = None,
    series: dict | None = None,
    series_start: datetime | None = None,
    step_minutes: int = 5,
    checks: list[dict] | None = None,
    meta: dict | None = None,
) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _summary_sheet(wb, report_label, site_data, meta or {})
    _episodes_sheet(wb, site_data)
    if series and series_start:
        _hourly_sheet(wb, site_data, series, series_start, step_minutes)
    _graphs_sheet(wb, screenshots_dir)
    if checks:
        _checks_sheet(wb, checks)
    path = output_dir / f"Utilization_Report_{report_label}.xlsx"
    wb.save(path)
    return path
